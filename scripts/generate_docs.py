"""Generate two project documentation DOCX files."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT_DIR = Path(__file__).resolve().parents[1] / "docs"


def setup_doc(title: str) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    for i in range(1, 4):
        hs = doc.styles[f"Heading {i}"]
        hs.font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)
        hs.font.name = "Calibri"
        hs._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    h = doc.add_heading(title, level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return doc


def add_para(doc: Document, text: str, *, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    p.paragraph_format.space_after = Pt(8)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            table.rows[ri + 1].cells[ci].text = val
    doc.add_paragraph()


def build_architecture() -> Path:
    doc = setup_doc("Архитектура и техническое описание")
    add_para(
        doc,
        "Модуль Organization Users — multi-tenant B2B-решение для управления пользователями "
        "организации: список, поиск, фильтры, добавление, смена роли/филиала, удаление из организации.",
    )

    doc.add_heading("1. Цель модуля", level=1)
    add_para(
        doc,
        "Реализовать изолированное по организациям управление пользователями с ролевой моделью (RBAC), "
        "без утечки факта существования чужих организаций и с понятными границами прав на чтение и изменение.",
    )

    doc.add_heading("2. Технологический стек", level=1)
    add_table(
        doc,
        ["Слой", "Технология", "Назначение"],
        [
            ["API", "FastAPI + SQLAlchemy 2 + Alembic", "REST API, ORM, миграции схем"],
            ["Auth", "JWT (Bearer) + org-scoped RBAC", "Аутентификация и авторизация"],
            ["БД", "PostgreSQL 16", "Хранение пользователей, орг., memberships"],
            ["UI", "React 18 + TypeScript + Vite", "SPA интерфейс"],
            ["Состояние UI", "TanStack Query", "Кэш и запросы к API"],
            ["i18n", "react-i18next (ru / kk / en)", "Локализация интерфейса"],
            ["Стили", "Tailwind CSS v4 + React Icons", "Адаптивный UI"],
            ["Runtime", "Docker Compose + nginx", "Локальный запуск полного стека"],
        ],
    )
    add_para(
        doc,
        "Почему FastAPI, а не Django: модуль API-first с явными границами авторизации. "
        "FastAPI даёт компактный код, OpenAPI-документацию и dependency injection для tenant/permission checks.",
    )

    doc.add_heading("3. Высокоуровневая архитектура", level=1)
    add_para(doc, "Компоненты при запуске через Docker Compose:")
    add_bullets(
        doc,
        [
            "nginx (порт хоста 8080) — единая точка входа: UI и прокси /api → backend",
            "frontend (Vite) — React SPA, внутри сети контейнеров",
            "backend (Uvicorn/FastAPI) — бизнес-логика и JWT",
            "PostgreSQL — персистентное хранилище",
        ],
    )
    add_para(
        doc,
        "Схема запросов: Браузер → http://localhost:8080 → nginx → / (frontend) или /api/… (backend). "
        "Фронтенд обращается к API по относительному пути /api, поэтому CORS и один origin упрощают демо.",
    )

    doc.add_heading("4. Multi-tenancy", level=1)
    add_para(
        doc,
        "Модель: shared database + row-level isolation через таблицу organization_memberships.",
        bold=False,
    )
    add_bullets(
        doc,
        [
            "User — глобальный аккаунт (один логин может состоять в нескольких организациях).",
            "Роль не глобальна: она привязана к membership (пара User ↔ Organization).",
            "Филиал (Branch) опционален и всегда принадлежит конкретной организации.",
            "Каждый маршрут /organizations/{organizationId}/… проходит через get_org_auth.",
        ],
    )
    add_para(doc, "Правила ответа get_org_auth:")
    add_bullets(
        doc,
        [
            "Нет активного membership → HTTP 404 (Organization not found) — без утечки существования org.",
            "Есть membership, но нет permission → HTTP 403.",
            "Невалидный/отсутствующий JWT → HTTP 401.",
        ],
    )

    doc.add_heading("5. Модель данных", level=1)
    add_para(doc, "Ключевые сущности:")
    add_bullets(
        doc,
        [
            "User — email (unique), password_hash, full_name, phone, is_active",
            "Organization — name, slug",
            "Branch — organization_id, name, is_active",
            "Role / Permission / RolePermission — RBAC-справочники",
            "OrganizationMembership — user_id, organization_id, role_id, branch_id?, is_active, joined_at",
        ],
    )
    add_para(doc, "Связи (упрощённо):")
    add_para(
        doc,
        "User ──< OrganizationMembership >── Organization\n"
        "                 │                       │\n"
        "                 ├── Role (+ Permissions)│\n"
        "                 └── Branch (optional) ──┘",
    )
    add_para(doc, "Ограничения:")
    add_bullets(
        doc,
        [
            "UNIQUE(user_id, organization_id) на memberships",
            "UNIQUE(organization_id, name) на branches",
            "FK membership.role → RESTRICT; membership.branch → SET NULL",
            "Удаление User/Organization каскадно удаляет memberships",
        ],
    )

    doc.add_heading("6. RBAC", level=1)
    add_table(
        doc,
        ["Permission", "Администратор", "Менеджер", "Оператор (employee)"],
        [
            ["users:read — просмотр списка/карточки", "да", "да", "да"],
            ["users:manage — создание/изменение/удаление", "да", "да", "нет"],
        ],
    )
    add_para(
        doc,
        "Проверки выполняются в FastAPI dependencies (get_org_auth + require_permission) "
        "и в сервисном слое (бизнес-правила).",
    )

    doc.add_heading("7. API", level=1)
    add_table(
        doc,
        ["Метод", "Путь", "Доступ"],
        [
            ["POST", "/auth/login", "публичный"],
            ["GET", "/auth/me", "аутентифицированный"],
            ["GET", "/organizations/{id}/users", "users:read"],
            ["POST", "/organizations/{id}/users", "users:manage"],
            ["PATCH", "/organizations/{id}/users/{userId}", "users:manage"],
            ["DELETE", "/organizations/{id}/users/{userId}", "users:manage"],
            ["GET", "/organizations/{id}/roles", "users:read"],
            ["GET", "/organizations/{id}/branches", "users:read"],
            ["GET", "/health", "публичный"],
        ],
    )
    add_para(
        doc,
        "Параметры списка пользователей: page, page_size, search, role_id, branch_id, status (active|inactive), "
        "sort_by (full_name|email|joined_at), sort_dir (asc|desc).",
    )
    add_para(
        doc,
        "Через nginx API доступен как http://localhost:8080/api/…. "
        "Интерактивная OpenAPI-документация backend (если порт проброшен): /docs.",
    )

    doc.add_heading("8. Структура backend", level=1)
    add_bullets(
        doc,
        [
            "app/main.py — приложение FastAPI, CORS, подключение роутеров",
            "app/api/ — HTTP-эндпоинты (auth, org_users)",
            "app/services/ — бизнес-логика пользователей организации",
            "app/auth/ — JWT, хеш паролей (bcrypt), deps (get_current_user, get_org_auth)",
            "app/models/ — SQLAlchemy-модели",
            "app/schemas/ — Pydantic DTO",
            "app/seed.py — идемпотентное наполнение демо-данными (~40 пользователей)",
            "alembic/ — миграции БД",
        ],
    )

    doc.add_heading("9. Структура frontend", level=1)
    add_bullets(
        doc,
        [
            "src/router.tsx — маршруты React Router",
            "src/routes/ProtectedRoute.tsx — редирект неавторизованных на /login",
            "src/auth/AuthContext.tsx — сессия, активная организация, permissions",
            "src/pages/LoginPage.tsx — вход",
            "src/pages/OrganizationUsersPage.tsx — основной экран модуля",
            "src/components/AppLayout.tsx — сайдбар, шапка, org switcher, i18n, logout",
            "src/api/client.ts — HTTP-клиент к /api",
            "src/i18n/ — локали ru / kk / en",
        ],
    )
    add_para(doc, "Маршруты UI:")
    add_bullets(
        doc,
        [
            "/login — авторизация",
            "/users — пользователи организации (основной функционал)",
            "/home, /organizations, /branches, /roles, /settings — заглушки навигации",
            "/ → редирект на /users",
        ],
    )

    doc.add_heading("10. Бизнес-правила и защита", level=1)
    add_bullets(
        doc,
        [
            "Добавление по email: если аккаунт есть — создаётся/реактивируется membership; если нет — нужен пароль.",
            "Удаление из org — soft-remove (membership.is_active = false), аккаунт сохраняется для других org.",
            "Нельзя удалить себя из организации.",
            "Нельзя деактивировать свой аккаунт.",
            "Нельзя понизить/удалить последнего активного администратора организации.",
            "Деактивация аккаунта глобальна (User.is_active): если пользователь состоит в других org — операция блокируется (предложить remove из текущей org).",
            "Филиал должен принадлежать той же организации.",
            "После деактивации аккаунта JWT перестаёт работать (проверка is_active на каждом запросе).",
        ],
    )

    doc.add_heading("11. Запуск (одна команда)", level=1)
    add_para(doc, "Единственное требование на машине ревьюера: установленный и запущенный Docker Desktop.")
    add_para(doc, "Команда:", bold=True)
    add_para(doc, "docker compose up --build")
    add_para(
        doc,
        "Эта команда сама: соберёт образы (все зависимости Python/Node внутри), поднимет PostgreSQL, "
        "выполнит миграции Alembic, засеет ~40 демо-пользователей и откроет приложение через nginx. "
        "Локально ставить Python, Node или Postgres не нужно. Файл .env создавать не нужно.",
    )
    add_para(doc, "Адрес приложения: http://localhost:8080", bold=True)
    add_table(
        doc,
        ["Сервис", "Адрес / порт"],
        [
            ["Приложение (nginx → UI + API)", "http://localhost:8080"],
            ["API через nginx", "http://localhost:8080/api/…"],
            ["PostgreSQL", "только внутри Docker-сети (порт на хост не публикуется)"],
        ],
    )
    add_para(
        doc,
        "Полный сброс БД и повторный seed: docker compose down -v && docker compose up --build",
    )
    add_para(
        doc,
        "Примечание: на Windows порт 80 часто занят IIS, поэтому наружный порт nginx — 8080.",
    )

    doc.add_heading("12. Безопасность (кратко)", level=1)
    add_bullets(
        doc,
        [
            "Пароли хранятся как bcrypt-хеш.",
            "Доступ к API — Bearer JWT.",
            "Tenant isolation на каждом org-scoped эндпоинте.",
            "Валидация входных данных через Pydantic (EmailStr, длины, enum-паттерны query).",
            "Демо JWT_SECRET в compose — только для локальной среды, не для production.",
        ],
    )

    path = OUT_DIR / "01_Architecture_and_Technical_Overview.docx"
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


def build_user_guide() -> Path:
    doc = setup_doc("Инструкция по пользованию")
    add_para(
        doc,
        "Документ описывает, как запустить модуль, войти в систему и проверить весь функционал "
        "управления пользователями организации без живой демонстрации.",
    )

    doc.add_heading("1. Как открыть приложение", level=1)
    add_para(doc, "Единственное требование: Docker Desktop (установлен и запущен).")
    add_para(doc, "Дополнительно ставить Python, Node.js, npm или PostgreSQL не нужно.")
    add_para(doc, "В корне проекта выполните одну команду:")
    add_para(doc, "docker compose up --build", bold=True)
    add_para(
        doc,
        "При первом запуске Docker скачает образы, установит зависимости внутри контейнеров, "
        "создаст БД, применит миграции и заполнит демо-данными (~40 пользователей).",
    )
    add_para(doc, "Откройте в браузере:")
    add_para(doc, "http://localhost:8080", bold=True)
    add_para(
        doc,
        "Подождите, пока в логах backend появится «Seed completed» и «Application startup complete», "
        "а nginx станет доступен (обычно до 1–2 минут на первом build). Затем обновите вкладку.",
    )
    add_para(doc, "Остановка: Ctrl+C в терминале или docker compose down.")
    add_para(doc, "Полный сброс данных: docker compose down -v && docker compose up --build")

    doc.add_heading("2. Авторизационные данные для теста", level=1)
    add_para(doc, "Пароль для всех демо-аккаунтов:", bold=True)
    add_para(doc, "password123", bold=True)
    add_para(doc, "Рекомендуемые учётки для проверки сценариев:")
    add_table(
        doc,
        ["Email", "Роль / организации", "Что демонстрирует"],
        [
            [
                "ivan@example.com",
                "Admin — ТОО Alpha Soft\nManager — Company B\nОператор — Company C",
                "Основной аккаунт. Переключение организаций меняет права. "
                "В Alpha Soft — полный CRUD и большой список (~40 users). "
                "В Company C — только просмотр.",
            ],
            [
                "anna@example.com",
                "Admin — Company B",
                "Админ одной организации. Чужие org в селекторе отсутствуют.",
            ],
            [
                "olga@example.com",
                "Manager — Alpha Soft\nОператор — Company B",
                "Управление в одной org и read-only в другой.",
            ],
            [
                "sergey@example.com",
                "Admin — Company C (часто единственный)",
                "Сценарий last-admin: нельзя снять/понизить последнего админа.",
            ],
            [
                "petr@example.com",
                "Оператор — Alpha Soft (аккаунт неактивен)",
                "Негативный сценарий: вход должен быть отклонён.",
            ],
        ],
    )
    add_para(
        doc,
        "В базе также есть десятки дополнительных пользователей (*@example.com) для пагинации, поиска и фильтров. "
        "Пароль у них тот же: password123.",
    )

    doc.add_heading("3. Вход в систему", level=1)
    add_bullets(
        doc,
        [
            "Откройте http://localhost:8080 — при отсутствии сессии откроется /login.",
            "Введите email и пароль, нажмите «Войти».",
            "Язык интерфейса (Рус / Қаз / Eng) можно сменить на экране входа и внутри приложения.",
            "После успеха откроется раздел «Пользователи» (/users).",
            "Кнопка выхода — красная иконка в шапке (и на мобильной версии).",
        ],
    )

    doc.add_heading("4. Интерфейс после входа", level=1)
    doc.add_heading("4.1. Навигация", level=2)
    add_bullets(
        doc,
        [
            "Слева (на desktop) / снизу (на mobile) — иконки разделов.",
            "Рабочий раздел модуля — «Пользователи» (иконка людей).",
            "Остальные пункты (Главная, Организации, Филиалы, Роли, Настройки) — заглушки «в разработке».",
        ],
    )
    doc.add_heading("4.2. Шапка", level=2)
    add_bullets(
        doc,
        [
            "Хлебные крошки / название текущего раздела.",
            "Селектор организации — переключает активный tenant и права.",
            "Переключатель языка.",
            "Колокольчик — демо-уведомления (без backend).",
            "Аватар / ФИО / роль в текущей организации.",
            "Иконка выхода.",
        ],
    )

    doc.add_heading("5. Раздел «Пользователи организации»", level=1)
    add_para(
        doc,
        "Экран показывает членов текущей выбранной организации. Данные загружаются с API с учётом прав.",
    )

    doc.add_heading("5.1. Просмотр списка", level=2)
    add_bullets(
        doc,
        [
            "Таблица: пользователь, email/телефон, роль, филиал, статус, дата добавления.",
            "На узких экранах — карточный вид.",
            "Пагинация внизу списка.",
            "Требуется permission users:read. Без него показывается сообщение об отсутствии прав.",
        ],
    )

    doc.add_heading("5.2. Поиск и фильтры", level=2)
    add_bullets(
        doc,
        [
            "Поиск по имени, email, телефону.",
            "Фильтр по роли.",
            "Фильтр по филиалу.",
            "Фильтр по статусу (Активен / Неактивен).",
            "Кнопка «Сбросить» очищает все фильтры.",
        ],
    )

    doc.add_heading("5.3. Добавление пользователя", level=2)
    add_para(doc, "Доступно при наличии users:manage (кнопка «+ Добавить пользователя»).")
    add_bullets(
        doc,
        [
            "Заполните ФИО, email, опционально телефон.",
            "Пароль: обязателен только если такого email ещё нет в системе.",
            "Если пользователь уже существует — он будет привязан к организации без нового пароля.",
            "Выберите роль и (опционально) филиал.",
            "Чекбокс «Отправить приглашение» — UI-демо, реальной почты нет.",
        ],
    )

    doc.add_heading("5.4. Просмотр и редактирование", level=2)
    add_bullets(
        doc,
        [
            "Меню «⋯» у строки → «Открыть» (карточка) или «Редактировать».",
            "В редактировании можно менять ФИО, телефон, роль, филиал, статус.",
            "Email при редактировании не меняется (идентификатор аккаунта).",
            "Свой статус деактивировать нельзя (поле блокируется / API вернёт ошибку).",
        ],
    )

    doc.add_heading("5.5. Удаление из организации", level=2)
    add_bullets(
        doc,
        [
            "Меню «⋯» → «Удалить» (нужен users:manage).",
            "Это soft-remove: пользователь исчезает из списка текущей org, но аккаунт может остаться в других.",
            "Удалить самого себя из UI нельзя (кнопка скрыта); API тоже запрещает.",
            "Нельзя удалить последнего активного администратора организации.",
        ],
    )

    doc.add_heading("6. Сценарии для самостоятельной проверки", level=1)

    doc.add_heading("Сценарий A. Полный доступ (Admin)", level=2)
    add_bullets(
        doc,
        [
            "Войти: ivan@example.com / password123.",
            "Выбрать организацию «ТОО Alpha Soft».",
            "Убедиться: виден большой список, работают поиск/фильтры/пагинация.",
            "Добавить тестового пользователя, отредактировать роль/филиал, затем удалить из организации.",
        ],
    )

    doc.add_heading("Сценарий B. Смена прав при смене организации", level=2)
    add_bullets(
        doc,
        [
            "Оставаясь под ivan@example.com, переключить org на «Company C».",
            "Кнопка добавления и пункты «Редактировать/Удалить» должны исчезнуть (роль Оператор).",
            "Список при этом доступен на чтение.",
            "Переключить на «Company B» — снова появляется управление (роль Менеджер).",
        ],
    )

    doc.add_heading("Сценарий C. Админ другой организации", level=2)
    add_bullets(
        doc,
        [
            "Выйти и войти как anna@example.com / password123.",
            "В селекторе только Company B.",
            "Полный CRUD внутри Company B.",
        ],
    )

    doc.add_heading("Сценарий D. Неактивный пользователь", level=2)
    add_bullets(
        doc,
        [
            "Попытаться войти как petr@example.com / password123.",
            "Ожидаемый результат: ошибка входа (аккаунт неактивен).",
        ],
    )

    doc.add_heading("Сценарий E. Last-admin", level=2)
    add_bullets(
        doc,
        [
            "Войти как sergey@example.com / password123 (Company C).",
            "Попытаться понизить себе роль до Оператора или деактивировать себя.",
            "Ожидается отказ с понятным сообщением об ошибке.",
        ],
    )

    doc.add_heading("Сценарий F. Мультитенантность UI", level=2)
    add_bullets(
        doc,
        [
            "Войти как olga@example.com.",
            "В Alpha Soft — права менеджера (можно управлять).",
            "В Company B — оператор (только просмотр).",
        ],
    )

    doc.add_heading("7. Языки интерфейса", level=1)
    add_para(
        doc,
        "Поддерживаются русский, қазақша и English. Переключатель в шапке и на экране логина. "
        "Содержимое данных (имена, названия org из seed) может оставаться на языке заполнения БД.",
    )

    doc.add_heading("8. Типичные сообщения об ошибках", level=1)
    add_table(
        doc,
        ["Ситуация", "Ожидаемое поведение"],
        [
            ["Неверный пароль", "Ошибка входа"],
            ["Неактивный аккаунт", "Ошибка входа"],
            ["Нет права manage", "Нет кнопок управления / API 403"],
            ["Удаление себя", "Запрещено"],
            ["Деактивация себя", "Запрещено"],
            ["Последний admin", "Нельзя понизить/удалить/деактивировать"],
            ["Филиал чужой org", "Ошибка валидации"],
            ["Пользователь уже в org", "Конфликт (уже участник)"],
        ],
    )

    doc.add_heading("9. Что не входит в модуль (осознанно)", level=1)
    add_bullets(
        doc,
        [
            "Реальная отправка email-приглашений.",
            "Полноценные экраны Филиалов / Ролей / Настроек (пока заглушки).",
            "Production-hardening JWT secret, TLS, отдельный CI/CD.",
        ],
    )

    doc.add_heading("10. Краткая шпаргалка", level=1)
    add_para(doc, "URL: http://localhost:8080")
    add_para(doc, "Главный логин: ivan@example.com")
    add_para(doc, "Пароль: password123")
    add_para(doc, "Команда запуска: docker compose up --build")

    path = OUT_DIR / "02_User_Guide_and_Test_Credentials.docx"
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


if __name__ == "__main__":
    # remove previous variants
    if OUT_DIR.exists():
        for old in OUT_DIR.glob("*.docx"):
            old.unlink(missing_ok=True)
    a = build_architecture()
    b = build_user_guide()
    print(a)
    print(b)
